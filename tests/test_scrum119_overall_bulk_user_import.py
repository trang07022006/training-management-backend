"""
Bộ kiểm thử tổng thể cho SCRUM-119: Bulk User Import - Activity A6.4.
Kiểm tra toàn bộ luồng Acceptance Criteria, Integration và Regression:
- Template Download: Header chuẩn, định dạng .xlsx
- Preview API: Đọc và kiểm tra dữ liệu từng dòng
- Execute Import API: Import partial success, invariants tổng số dòng
- Validation: Name, email, phone, role, status, password
- Duplicate Email: Trùng lặp trong DB và trùng lặp trong cùng file
- File Format & Integrity: File rỗng, file hỏng, thiếu sheet Users, sai header
- Phân quyền (RBAC): Chỉ user có quyền user.manage mới được thao tác
- Bảo mật (Security): Không để lộ password plaintext hoặc password hash trong response
"""

from datetime import datetime, timedelta, timezone
from io import BytesIO
import unittest
from unittest.mock import MagicMock, patch
from flask import Flask
from flask_jwt_extended import JWTManager, create_access_token
import openpyxl

from app.models.login_session import LoginSession
from app.models.permission import Permission
from app.models.role import Role
from app.routes.user_routes import user_bp
from app.services.user_import_service import USER_IMPORT_TEMPLATE_HEADERS
from app.services.user_service import DuplicateEmailError


def _create_excel(rows=None, sheet_name="Users", headers=None) -> BytesIO:
    """Helper tạo workbook in-memory phục vụ kiểm thử."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers if headers is not None else USER_IMPORT_TEMPLATE_HEADERS)
    if rows:
        for r in rows:
            ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class Scrum119OverallBulkUserImportTests(unittest.TestCase):
    """Test suite tổng thể cho SCRUM-119 Activity A6.4."""

    def setUp(self):
        self.app = Flask(__name__)
        self.app.config.update(
            TESTING=True,
            JWT_SECRET_KEY="test-secret-key-scrum119-overall-suite-32-bytes",
            JWT_ACCESS_TOKEN_EXPIRES=timedelta(minutes=15),
            SESSION_IDLE_TIMEOUT=timedelta(minutes=30),
        )
        self.jwt = JWTManager(self.app)
        self.app.register_blueprint(user_bp)

        self.client = self.app.test_client()
        self.now = datetime.now(timezone.utc)
        self.user_id = "admin-scrum119-overall"
        self.session_id = "session-scrum119-overall"

        self.session = LoginSession(
            session_id=self.session_id,
            user_id=self.user_id,
            refresh_jti="ref-jti-119-overall",
            created_at=self.now,
            last_activity_at=self.now,
            idle_expires_at=self.now + timedelta(minutes=30),
            absolute_expires_at=self.now + timedelta(hours=8),
        )

    def _create_token(self, user_id=None, session_id=None):
        with self.app.app_context():
            return create_access_token(
                identity=user_id or self.user_id,
                additional_claims={"sid": session_id or self.session_id},
            )

    def _setup_rbac_admin(self, mock_get_session, mock_perm_repo, mock_ur_roles, mock_rp_codes):
        mock_get_session.return_value = self.session
        mock_perm_repo.return_value = Permission(id="p-manage", code="user.manage", name="Quản lý user")
        mock_ur_roles.return_value = [Role(id="r-admin", code="admin", name="Admin")]
        mock_rp_codes.return_value = ["user.manage"]

    def _setup_rbac_no_perm(self, mock_get_session, mock_perm_repo, mock_ur_roles, mock_rp_codes):
        mock_get_session.return_value = self.session
        mock_perm_repo.return_value = Permission(id="p-view", code="user.view", name="Xem user")
        mock_ur_roles.return_value = [Role(id="r-trainer", code="trainer", name="Trainer")]
        mock_rp_codes.return_value = ["user.view"]

    # CASE 1: Download Template
    @patch("app.utils.rbac_auth.get_session_by_id")
    @patch("app.services.authorization_service.permission_repository.get_permission_by_code")
    @patch("app.services.authorization_service.user_role_repository.get_roles_by_user_id")
    @patch("app.services.authorization_service.role_permission_repository.get_permission_codes_by_role_id")
    def test_case01_download_template_success(self, mock_rp, mock_ur, mock_p, mock_s):
        """CASE 1: Tải template chuẩn thành công với header name, email, password, phone, role, status."""
        self._setup_rbac_admin(mock_s, mock_p, mock_ur, mock_rp)
        token = self._create_token()

        res = self.client.get(
            "/api/users/import/template",
            headers={"Authorization": f"Bearer {token}"},
        )
        self.assertEqual(res.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", res.content_type)

        wb = openpyxl.load_workbook(BytesIO(res.data))
        self.assertIn("Users", wb.sheetnames)
        ws = wb["Users"]
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers, list(USER_IMPORT_TEMPLATE_HEADERS))

    # CASE 2: File hợp lệ 100% (Preview -> Import)
    @patch("app.utils.rbac_auth.get_session_by_id")
    @patch("app.services.authorization_service.permission_repository.get_permission_by_code")
    @patch("app.services.authorization_service.user_role_repository.get_roles_by_user_id")
    @patch("app.services.authorization_service.role_permission_repository.get_permission_codes_by_role_id")
    @patch("app.repositories.user_repository.check_email_exists")
    @patch("app.services.user_service.create_user_service")
    def test_case02_valid_100_percent_import(
        self,
        mock_create_user,
        mock_check_email,
        mock_rp,
        mock_ur,
        mock_p,
        mock_s,
    ):
        """CASE 2: File hợp lệ 100% -> preview valid_rows=total_rows, import imported_rows=total_rows."""
        self._setup_rbac_admin(mock_s, mock_p, mock_ur, mock_rp)
        mock_check_email.return_value = False

        rows = [
            ["Nguyễn Văn A", "a@daotao.edu.vn", "Password123", "0912345678", "trainee", "active"],
            ["Trần Thị B", "b@daotao.edu.vn", "Password456", "0987654321", "trainer", "active"],
        ]
        buf = _create_excel(rows=rows)
        token = self._create_token()

        # Preview
        preview_res = self.client.post(
            "/api/users/import/preview",
            headers={"Authorization": f"Bearer {token}"},
            data={"file": (BytesIO(buf.getvalue()), "users.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(preview_res.status_code, 200)
        p_data = preview_res.get_json()["data"]
        self.assertEqual(p_data["total_rows"], 2)
        self.assertEqual(p_data["valid_rows"], 2)
        self.assertEqual(p_data["invalid_rows"], 0)

        # Import
        import_res = self.client.post(
            "/api/users/import",
            headers={"Authorization": f"Bearer {token}"},
            data={"file": (BytesIO(buf.getvalue()), "users.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(import_res.status_code, 200)
        i_data = import_res.get_json()["data"]
        self.assertEqual(i_data["total_rows"], 2)
        self.assertEqual(i_data["imported_rows"], 2)
        self.assertEqual(i_data["skipped_rows"], 0)
        self.assertEqual(mock_create_user.call_count, 2)

    # CASE 3 & 11: File có dòng lỗi validation & Partial success
    @patch("app.utils.rbac_auth.get_session_by_id")
    @patch("app.services.authorization_service.permission_repository.get_permission_by_code")
    @patch("app.services.authorization_service.user_role_repository.get_roles_by_user_id")
    @patch("app.services.authorization_service.role_permission_repository.get_permission_codes_by_role_id")
    @patch("app.repositories.user_repository.check_email_exists")
    @patch("app.services.user_service.create_user_service")
    def test_case03_and_11_partial_success_with_validation_errors(
        self,
        mock_create_user,
        mock_check_email,
        mock_rp,
        mock_ur,
        mock_p,
        mock_s,
    ):
        """CASE 3 & 11: File hỗn hợp (1 hợp lệ, 2 lỗi validation) -> chỉ import 1 dòng, bỏ qua 2 dòng lỗi."""
        self._setup_rbac_admin(mock_s, mock_p, mock_ur, mock_rp)
        mock_check_email.return_value = False

        rows = [
            ["Nguyễn Hợp Lệ", "valid@daotao.edu.vn", "Password123", "0912345678", "trainee", "active"],
            ["", "no_name@daotao.edu.vn", "Password123", "0912345678", "trainee", "active"],
            ["Email Sai", "not-an-email", "Password123", "0912345678", "trainee", "active"],
        ]
        buf = _create_excel(rows=rows)
        token = self._create_token()

        res = self.client.post(
            "/api/users/import",
            headers={"Authorization": f"Bearer {token}"},
            data={"file": (buf, "mixed.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res.status_code, 200)
        data = res.get_json()["data"]

        # Invariants: total_rows = imported_rows + skipped_rows
        self.assertEqual(data["total_rows"], 3)
        self.assertEqual(data["imported_rows"], 1)
        self.assertEqual(data["skipped_rows"], 2)
        self.assertEqual(data["summary"]["validation_errors"], 2)
        self.assertEqual(mock_create_user.call_count, 1)

    # CASE 4: Duplicate Email trong Database
    @patch("app.utils.rbac_auth.get_session_by_id")
    @patch("app.services.authorization_service.permission_repository.get_permission_by_code")
    @patch("app.services.authorization_service.user_role_repository.get_roles_by_user_id")
    @patch("app.services.authorization_service.role_permission_repository.get_permission_codes_by_role_id")
    @patch("app.repositories.user_repository.check_email_exists")
    @patch("app.services.user_service.create_user_service")
    def test_case04_duplicate_email_in_database(
        self,
        mock_create_user,
        mock_check_email,
        mock_rp,
        mock_ur,
        mock_p,
        mock_s,
    ):
        """CASE 4: Email đã tồn tại trong DB -> bị bắt ở preview (DUPLICATE_EMAIL)."""
        self._setup_rbac_admin(mock_s, mock_p, mock_ur, mock_rp)
        mock_check_email.side_effect = lambda email: email == "exist@daotao.edu.vn"

        rows = [
            ["User Trùng", "exist@daotao.edu.vn", "Password123", "0912345678", "trainee", "active"],
        ]
        buf = _create_excel(rows=rows)
        token = self._create_token()

        res = self.client.post(
            "/api/users/import/preview",
            headers={"Authorization": f"Bearer {token}"},
            data={"file": (buf, "dup_db.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res.status_code, 200)
        row = res.get_json()["data"]["rows"][0]
        self.assertFalse(row["valid"])
        error_codes = [e["code"] for e in row["errors"]]
        self.assertIn("DUPLICATE_EMAIL", error_codes)

    # CASE 5: Duplicate Email trong cùng 1 file
    @patch("app.utils.rbac_auth.get_session_by_id")
    @patch("app.services.authorization_service.permission_repository.get_permission_by_code")
    @patch("app.services.authorization_service.user_role_repository.get_roles_by_user_id")
    @patch("app.services.authorization_service.role_permission_repository.get_permission_codes_by_role_id")
    @patch("app.repositories.user_repository.check_email_exists")
    def test_case05_duplicate_email_within_same_file(self, mock_check_email, mock_rp, mock_ur, mock_p, mock_s):
        """CASE 5: Hai dòng có cùng email trong file -> dòng sau bị lỗi DUPLICATE_EMAIL_IN_FILE."""
        self._setup_rbac_admin(mock_s, mock_p, mock_ur, mock_rp)
        mock_check_email.return_value = False

        rows = [
            ["Dòng 1", "same@daotao.edu.vn", "Password123", "0912345678", "trainee", "active"],
            ["Dòng 2", "same@daotao.edu.vn", "Password456", "0987654321", "trainer", "active"],
        ]
        buf = _create_excel(rows=rows)
        token = self._create_token()

        res = self.client.post(
            "/api/users/import/preview",
            headers={"Authorization": f"Bearer {token}"},
            data={"file": (buf, "dup_file.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res.status_code, 200)
        rows_data = res.get_json()["data"]["rows"]
        self.assertTrue(rows_data[0]["valid"])
        self.assertFalse(rows_data[1]["valid"])
        self.assertEqual(rows_data[1]["errors"][0]["code"], "DUPLICATE_EMAIL_IN_FILE")

    # CASE 6 & 7: File sai định dạng hoặc rỗng
    @patch("app.utils.rbac_auth.get_session_by_id")
    @patch("app.services.authorization_service.permission_repository.get_permission_by_code")
    @patch("app.services.authorization_service.user_role_repository.get_roles_by_user_id")
    @patch("app.services.authorization_service.role_permission_repository.get_permission_codes_by_role_id")
    def test_case06_and_07_invalid_extension_and_empty(self, mock_rp, mock_ur, mock_p, mock_s):
        """CASE 6 & 7: Không phải .xlsx trả INVALID_FILE_TYPE; file không byte trả FILE_EMPTY."""
        self._setup_rbac_admin(mock_s, mock_p, mock_ur, mock_rp)
        token = self._create_token()

        # Sai định dạng
        res_csv = self.client.post(
            "/api/users/import/preview",
            headers={"Authorization": f"Bearer {token}"},
            data={"file": (BytesIO(b"name,email"), "users.csv")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res_csv.status_code, 400)
        self.assertEqual(res_csv.get_json()["code"], "INVALID_FILE_TYPE")

        # File rỗng 0 bytes
        res_empty = self.client.post(
            "/api/users/import/preview",
            headers={"Authorization": f"Bearer {token}"},
            data={"file": (BytesIO(b""), "users.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res_empty.status_code, 400)
        self.assertEqual(res_empty.get_json()["code"], "FILE_EMPTY")

    # CASE 8, 9, 10: File hỏng, thiếu sheet Users, sai header
    @patch("app.utils.rbac_auth.get_session_by_id")
    @patch("app.services.authorization_service.permission_repository.get_permission_by_code")
    @patch("app.services.authorization_service.user_role_repository.get_roles_by_user_id")
    @patch("app.services.authorization_service.role_permission_repository.get_permission_codes_by_role_id")
    def test_case08_09_10_corrupt_missing_sheet_invalid_header(self, mock_rp, mock_ur, mock_p, mock_s):
        """CASE 8, 9, 10: File hỏng -> INVALID_EXCEL_FILE, thiếu sheet -> MISSING_USERS_SHEET, sai header -> INVALID_HEADER."""
        self._setup_rbac_admin(mock_s, mock_p, mock_ur, mock_rp)
        token = self._create_token()

        # File hỏng
        res_corrupt = self.client.post(
            "/api/users/import/preview",
            headers={"Authorization": f"Bearer {token}"},
            data={"file": (BytesIO(b"not-a-valid-zip-xlsx"), "corrupt.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res_corrupt.status_code, 400)
        self.assertEqual(res_corrupt.get_json()["code"], "INVALID_EXCEL_FILE")

        # Thiếu sheet Users
        buf_no_sheet = _create_excel(sheet_name="WrongSheet")
        res_sheet = self.client.post(
            "/api/users/import/preview",
            headers={"Authorization": f"Bearer {token}"},
            data={"file": (buf_no_sheet, "nosheet.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res_sheet.status_code, 400)
        self.assertEqual(res_sheet.get_json()["code"], "MISSING_USERS_SHEET")

        # Sai header
        buf_bad_header = _create_excel(headers=["ho_ten", "thu_dien_tu"])
        res_header = self.client.post(
            "/api/users/import/preview",
            headers={"Authorization": f"Bearer {token}"},
            data={"file": (buf_bad_header, "bad_header.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res_header.status_code, 400)
        self.assertEqual(res_header.get_json()["code"], "INVALID_HEADER")

    # CASE 13: RBAC Protection
    @patch("app.utils.rbac_auth.get_session_by_id")
    @patch("app.services.authorization_service.permission_repository.get_permission_by_code")
    @patch("app.services.authorization_service.user_role_repository.get_roles_by_user_id")
    @patch("app.services.authorization_service.role_permission_repository.get_permission_codes_by_role_id")
    def test_case13_rbac_unauthorized_access_returns_403(self, mock_rp, mock_ur, mock_p, mock_s):
        """CASE 13: User không có quyền user.manage bị chặn 403 ở cả /template, /preview, /import."""
        self._setup_rbac_no_perm(mock_s, mock_p, mock_ur, mock_rp)
        token = self._create_token()

        res1 = self.client.get("/api/users/import/template", headers={"Authorization": f"Bearer {token}"})
        self.assertEqual(res1.status_code, 403)

        res2 = self.client.post("/api/users/import/preview", headers={"Authorization": f"Bearer {token}"})
        self.assertEqual(res2.status_code, 403)

        res3 = self.client.post("/api/users/import", headers={"Authorization": f"Bearer {token}"})
        self.assertEqual(res3.status_code, 403)

    # CASE 14: Security
    @patch("app.utils.rbac_auth.get_session_by_id")
    @patch("app.services.authorization_service.permission_repository.get_permission_by_code")
    @patch("app.services.authorization_service.user_role_repository.get_roles_by_user_id")
    @patch("app.services.authorization_service.role_permission_repository.get_permission_codes_by_role_id")
    @patch("app.repositories.user_repository.check_email_exists")
    @patch("app.services.user_service.create_user_service")
    def test_case14_security_no_plaintext_or_hashed_passwords(
        self,
        mock_create,
        mock_check_email,
        mock_rp,
        mock_ur,
        mock_p,
        mock_s,
    ):
        """CASE 14: Tuyệt đối không để lộ plaintext password hoặc hash trong response của preview lẫn import."""
        self._setup_rbac_admin(mock_s, mock_p, mock_ur, mock_rp)
        mock_check_email.return_value = False

        super_secret = "SuperSecretPassword123!"
        rows = [["Bảo Mật", "sec@daotao.edu.vn", super_secret, "0912345678", "trainee", "active"]]
        buf = _create_excel(rows=rows)
        token = self._create_token()

        # Check preview
        res_preview = self.client.post(
            "/api/users/import/preview",
            headers={"Authorization": f"Bearer {token}"},
            data={"file": (BytesIO(buf.getvalue()), "sec.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res_preview.status_code, 200)
        self.assertNotIn(super_secret, res_preview.get_data(as_text=True))
        self.assertNotIn("password_hash", res_preview.get_data(as_text=True))

        # Check import
        res_import = self.client.post(
            "/api/users/import",
            headers={"Authorization": f"Bearer {token}"},
            data={"file": (BytesIO(buf.getvalue()), "sec.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res_import.status_code, 200)
        self.assertNotIn(super_secret, res_import.get_data(as_text=True))
        self.assertNotIn("password_hash", res_import.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
