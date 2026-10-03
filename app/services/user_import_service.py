"""
Service hỗ trợ bulk import người dùng qua Excel (SCRUM-119).
Activity A1: Tạo file Excel mẫu cho bulk import người dùng.
"""

from io import BytesIO
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


USER_IMPORT_TEMPLATE_HEADERS = [
    "name",
    "email",
    "password",
    "phone",
    "role",
    "status",
]

USER_IMPORT_TEMPLATE_FILENAME = "user_import_template.xlsx"
EXCEL_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def generate_user_import_template() -> BytesIO:
    """
    Tạo file Excel mẫu (.xlsx) phục vụ import người dùng.
    - Sheet 1 ('Users'): chứa header chuẩn theo contract tạo user và 1 dòng dữ liệu mẫu an toàn.
    - Sheet 2 ('Huong_Dan'): mô tả chi tiết quy định từng trường, ràng buộc bắt buộc và giá trị hợp lệ.
    """
    wb = openpyxl.Workbook()

    # --- Sheet 1: Users ---
    ws_users = wb.active
    ws_users.title = "Users"

    # Header styling
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")

    ws_users.append(USER_IMPORT_TEMPLATE_HEADERS)
    for col_idx in range(1, len(USER_IMPORT_TEMPLATE_HEADERS) + 1):
        cell = ws_users.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Dữ liệu mẫu minh họa an toàn (không chứa mật khẩu thật hay email thật nhạy cảm)
    sample_row = [
        "Nguyen Van A",
        "example.user@daotao.edu.vn",
        "DemoUser@123",
        "0912345678",
        "trainee",
        "active",
    ]
    ws_users.append(sample_row)

    # Đặt độ rộng cột cho dễ nhìn
    column_widths = {
        "A": 25,  # name
        "B": 32,  # email
        "C": 18,  # password
        "D": 18,  # phone
        "E": 15,  # role
        "F": 15,  # status
    }
    for col_letter, width in column_widths.items():
        ws_users.column_dimensions[col_letter].width = width

    # --- Sheet 2: Huong_Dan ---
    ws_guide = wb.create_sheet(title="Huong_Dan")
    guide_headers = ["Tên cột (Header)", "Bắt buộc", "Kiểu dữ liệu", "Quy tắc và giá trị hợp lệ", "Mô tả"]
    ws_guide.append(guide_headers)

    guide_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    guide_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col_idx in range(1, len(guide_headers) + 1):
        cell = ws_guide.cell(row=1, column=col_idx)
        cell.fill = guide_fill
        cell.font = guide_font
        cell.alignment = center_align

    guide_data = [
        ("name", "Bắt buộc", "Chuỗi ký tự", "Không được để trống", "Họ và tên người dùng"),
        ("email", "Bắt buộc", "Email", "Định dạng email hợp lệ, tối đa 255 ký tự, không trùng lặp trong hệ thống", "Địa chỉ email đăng nhập"),
        ("password", "Bắt buộc", "Chuỗi ký tự", "Tối thiểu 8 ký tự, phải bao gồm cả chữ cái và chữ số", "Mật khẩu khởi tạo ban đầu"),
        ("phone", "Tùy chọn", "Số điện thoại", "9-15 chữ số (vd: 0912345678 hoặc +84912345678)", "Số điện thoại liên hệ"),
        ("role", "Tùy chọn", "Mã vai trò", "admin | trainer | trainee | user (mặc định: user)", "Vai trò chính của tài khoản"),
        ("status", "Tùy chọn", "Trạng thái", "active | inactive | locked (mặc định: active)", "Trạng thái kích hoạt của tài khoản"),
    ]
    for row in guide_data:
        ws_guide.append(list(row))

    guide_column_widths = {"A": 20, "B": 15, "C": 18, "D": 45, "E": 35}
    for col_letter, width in guide_column_widths.items():
        ws_guide.column_dimensions[col_letter].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
